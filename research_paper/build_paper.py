"""Build Paper Script
Generates both research_paper.docx and research_paper.pdf from the validated manuscript and experimental figures.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, HRFlowable
)

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RESEARCH_PAPER_DIR = os.path.join(PROJECT_ROOT, "research_paper")
FIGURES_DIR = os.path.join(PROJECT_ROOT, "paper", "figures")

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_docx():
    doc = Document()
    
    # Page setup - Normal 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("Physics-Constrained AI Agents for Automated Quantum Circuit Debugging and Repair")
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x08, 0x4B, 0x6E)

    # Author
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(14)
    r_author = p_author.add_run("Alok Pandey\n")
    r_author.font.size = Pt(12)
    r_author.font.bold = True
    r_affil = p_author.add_run("Department of Computer Science and Engineering\nIndian Institute of Technology Patna, India")
    r_affil.font.size = Pt(10.5)
    r_affil.font.italic = True

    # Abstract Box
    table_abs = doc.add_table(rows=1, cols=1)
    table_abs.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_abs = table_abs.cell(0, 0)
    cell_abs.width = Inches(6.5)
    set_cell_background(cell_abs, "F0F4F8")
    
    p_abs = cell_abs.paragraphs[0]
    p_abs.paragraph_format.space_after = Pt(4)
    r_abs_title = p_abs.add_run("Abstract—")
    r_abs_title.bold = True
    r_abs_text = p_abs.add_run(
        "Quantum-circuit debugging is difficult because a syntactically valid or logically plausible repair may still violate "
        "hardware-native gates, coupling connectivity, or noise-aware execution limits. This paper presents a physics-constrained "
        "agentic framework that couples fault localization and minimal candidate generation with compilation and explicit validation. "
        "A structured circuit state records diagnostics and repair history; failed checks provide feedback for subsequent attempts. "
        "The reported benchmark comprises 1,000 buggy OpenQASM 3.0 circuits across five bug classes and three compared systems. "
        "The proposed agent achieves Pass@1 of 77.5%, Pass@3 of 99.1%, and hardware physics validity of 99.2%, with a mean of 1.31 repair iterations. "
        "The evaluation grounds circuits against the IBM FakeSherbrooke 127-qubit superconducting processor model. "
        "Statistical tests and effect size analyses confirm that hardware-grounded validation substantially outperforms simulation-only and heuristic baselines."
    )
    r_abs_text.font.size = Pt(10)

    p_kw = cell_abs.add_paragraph()
    p_kw.paragraph_format.space_after = Pt(0)
    r_kw_title = p_kw.add_run("Keywords—")
    r_kw_title.bold = True
    r_kw_title.font.italic = True
    r_kw_text = p_kw.add_run("quantum software engineering, circuit debugging, AI agents, automated program repair, hardware-aware compilation, superconducting quantum processors.")
    r_kw_text.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_sec(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x08, 0x4B, 0x6E)
        return p

    def add_subsec(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return p

    # SECTION I
    add_sec("I. INTRODUCTION")
    doc.add_paragraph(
        "Quantum computing systems require precise agreement between high-level computational specifications and physically executable quantum circuits. "
        "A local textual modification can easily resolve a syntax error while erroneously modifying an entangling target, disrupting basis rotation angles, "
        "or generating multi-qubit interactions prohibited by device topology. Conversely, an algorithmically sound circuit typically requires rigorous "
        "decomposition and SWAP routing before it can execute on physical quantum hardware. Consequently, quantum program repair is intrinsically more demanding "
        "than classical code repair: patches cannot be judged solely by syntactic validity or unconstrained simulation."
    )
    doc.add_paragraph(
        "Prior language-model-based quantum repair tools (such as QBugLM) rely on unconstrained simulation, ignoring coupling graphs, gate infidelities, and coherence decay limits. "
        "Symbolic optimizers (e.g., PyZX) simplify ZX-calculus diagrams but lack defect-targeted repair mechanisms for algorithmic errors. "
        "To address this gap, we present an end-to-end physics-constrained multi-agent framework that integrates fault localization, minimal candidate synthesis, and rigorous hardware validation into a unified feedback loop. "
        "Grounding repairs on IBM's 127-qubit FakeSherbrooke Eagle-generation superconducting processor, our agent verifies two-qubit coupling graphs, calibrated gate error infidelities, pulse durations, and T1/T2 coherence limits."
    )

    # SECTION II
    add_sec("II. RELATED WORK AND HARDWARE CONSTRAINTS")
    add_subsec("A. Language-Guided Quantum Program Repair")
    doc.add_paragraph(
        "Initial explorations by Guo et al. established the feasibility of repairing quantum code using LLM zero-shot prompts. "
        "Recently, Pham et al. introduced QBugLM as an agentic benchmark. However, QBugLM evaluates candidate patches solely through ideal statevector simulation. "
        "When deployed to physical NISQ devices, simulation-valid circuits frequently fail due to prohibited two-qubit couplings or excessive circuit depth that exceeds qubit coherence lifetimes."
    )
    add_subsec("B. Symbolic Simplification and ZX-Calculus")
    doc.add_paragraph(
        "PyZX provides rigorous graph-theoretic circuit simplification via the ZX-calculus. While powerful for gate cancellation and T-count reduction, "
        "PyZX is inherently an optimizer, not an automated repair tool. It cannot synthesize missing operations, correct corrupted rotation parameters, or resolve parser syntax exceptions."
    )
    add_subsec("C. Physical Superconducting Hardware Constraints")
    doc.add_paragraph(
        "Superconducting quantum processors exhibit strict physical limitations. In IBM FakeSherbrooke (127 qubits): "
        "(1) Coupling topology follows a sparse heavy-hexagonal lattice where CNOT/ECR operations can only occur between directly connected physical qubits; "
        "(2) Physical gate executions require finite pulse durations (e.g., ~300 ns for entangling gates), accumulating operational time; "
        "(3) Qubits undergo thermal relaxation (T1 ≈ 250 μs) and dephasing (T2 ≈ 150 μs), setting a hard ceiling on viable circuit latency."
    )

    # SECTION III
    add_sec("III. PROBLEM FORMULATION & BENCHMARK TAXONOMY")
    doc.add_paragraph(
        "Let C_bug be a buggy OpenQASM 3.0 circuit, O an independent behavioral oracle (ground-truth unitary or statevector), and B a physical hardware backend. "
        "A repair system synthesizes a candidate C* satisfying four joint predicates:"
    )
    doc.add_paragraph(
        "    Accept(C*) = V_syntax(C*) ∧ V_structural(C*) ∧ V_semantic(C*, O) ∧ V_physics(C*, B)\n"
        "where V_physics verifies connectivity over the backend coupling graph G=(V,E), gate duration bounds, and coherence constraints."
    )

    # Table I in Word
    p_t1 = doc.add_paragraph()
    r_t1 = p_t1.add_run("TABLE I: BENCHMARK DEFECT TAXONOMY (N = 1,000 CIRCUITS, 200 PER CATEGORY)")
    r_t1.bold = True
    r_t1.font.size = Pt(10)
    
    t1 = doc.add_table(rows=6, cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1_headers = ["Bug Category", "Description & Injection Mechanism", "Validation Criteria"]
    for col_idx, h in enumerate(t1_headers):
        cell = t1.cell(0, col_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "084B6E")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    t1_data = [
        ("Syntax Error", "Malformed QASM 3.0 syntax, omitted semicolons, corrupt keywords, unclosed brackets.", "OpenQASM parser diagnostic validation"),
        ("Structural Error", "Missing entangling/superposition gates, inverted control/target qubit operands.", "AST gate sequence & register bounds checking"),
        ("Gate Redundancy", "Self-canceling duplicate gates (X·X, H·H, identical CX·CX consecutive pairs).", "Circuit depth & redundant operator cancellation"),
        ("Semantic Deviation", "Inverted rotation angles (-θ), shifted phase angles, basis mutations (Rx ↔ Rz).", "Statevector fidelity F = |⟨ψ_gt|ψ_rep⟩|² ≥ 0.999"),
        ("Physics Violation", "Non-nearest-neighbor coupling violations on FakeSherbrooke, delays exceeding T2.", "Heavy-hex connectivity graph & coherence check")
    ]
    for row_idx, row in enumerate(t1_data):
        for col_idx, text in enumerate(row):
            cell = t1.cell(row_idx + 1, col_idx)
            cell.text = text
            if row_idx % 2 == 1:
                set_cell_background(cell, "F9FAFB")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # SECTION IV
    add_sec("IV. MULTI-AGENT REPAIR ARCHITECTURE")
    doc.add_paragraph(
        "The proposed framework decomposes the repair problem into three coordinated agents orchestrated in an iterative feedback loop:"
    )
    doc.add_paragraph(
        "1. Diagnostician Agent: Ingests parser logs, statevector discrepancies, and physical validation reports to localize the precise line number, faulty qubit operands, and bug classification.\n"
        "2. Repairer Agent: Synthesizes minimal, targeted OpenQASM transformations (e.g., syntax correction, redundant pair elimination, parameter angle inversion, or nearest-neighbor SWAP insertion).\n"
        "3. Validator Agent: Executes the compiled candidate against both the mathematical semantic oracle and the FakeSherbrooke hardware physics model. Detailed diagnostic reports are passed back to the orchestrator upon failure."
    )

    # Insert Architecture / Workflow summary
    p_diag = doc.add_paragraph()
    r_diag = p_diag.add_run("Multi-Agent Coordination Cycle: ")
    r_diag.bold = True
    p_diag.add_run("Input C_bug → Diagnostician → Repairer Candidate C_t → Validator (Syntax → Semantics → Physics) → Acceptance or Failure Feedback (up to k=3 iterations).")

    # SECTION V & VI: RESULTS
    add_sec("V. EXPERIMENTAL BENCHMARK & RESULTS")
    doc.add_paragraph(
        "We evaluated the three systems across all 1,000 circuits in the benchmark. All experiments executed against the 127-qubit IBM FakeSherbrooke backend."
    )

    # Table II: Overall Performance
    p_t2 = doc.add_paragraph()
    r_t2 = p_t2.add_run("TABLE II: OVERALL REPAIR PERFORMANCE ACROSS 1,000 BENCHMARK CIRCUITS")
    r_t2.bold = True
    r_t2.font.size = Pt(10)

    t2 = doc.add_table(rows=4, cols=7)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2_headers = ["System", "Pass@1 (%)", "Pass@3 (%)", "Physics Valid (%)", "Gate Red.", "Depth Red.", "Mean Iters"]
    for col_idx, h in enumerate(t2_headers):
        cell = t2.cell(0, col_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "084B6E")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t2_data = [
        ("Physics-Constrained Agent (Ours)", "77.5 ± 41.8", "99.1 ± 9.4", "99.2 ± 8.9", "-9.5 ± 16.8", "-10.0 ± 17.2", "1.31 ± 0.62"),
        ("QBugLM (Unconstrained LLM)", "38.1 ± 48.6", "38.1 ± 48.6", "69.6 ± 46.0", "-7.8 ± 15.6", "-8.1 ± 16.0", "1.00 ± 0.00"),
        ("PyZX (ZX-Calculus Heuristic)", "5.0 ± 21.8", "5.0 ± 21.8", "15.5 ± 36.2", "-8.4 ± 13.8", "-6.3 ± 11.0", "1.00 ± 0.00")
    ]
    for row_idx, row in enumerate(t2_data):
        for col_idx, text in enumerate(row):
            cell = t2.cell(row_idx + 1, col_idx)
            cell.text = text
            if row_idx == 0:
                set_cell_background(cell, "E8F4F8")
                cell.paragraphs[0].runs[0].font.bold = True
            elif row_idx % 2 == 1:
                set_cell_background(cell, "F9FAFB")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Insert Figures if present
    fig_pass = os.path.join(FIGURES_DIR, "pass_at_k_comparison.png")
    if os.path.exists(fig_pass):
        p_fig1 = doc.add_paragraph()
        p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(fig_pass, width=Inches(5.5))
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c1 = p_cap1.add_run("Fig. 1. Pass@1 and Pass@3 repair success rates compared across systems.")
        r_c1.font.italic = True
        r_c1.font.size = Pt(9.5)

    # Table III: Defect Category Breakdown
    p_t3 = doc.add_paragraph()
    r_t3 = p_t3.add_run("TABLE III: REPAIR SUCCESS (PASS@3 %) AND PHYSICS VALIDITY (%) BY DEFECT CATEGORY (N=200 PER CATEGORY)")
    r_t3.bold = True
    r_t3.font.size = Pt(10)

    t3 = doc.add_table(rows=6, cols=7)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3_headers = ["Bug Category", "PyZX Pass@3", "QBugLM Pass@3", "Ours Pass@3", "PyZX Phys.", "QBugLM Phys.", "Ours Phys."]
    for col_idx, h in enumerate(t3_headers):
        cell = t3.cell(0, col_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "084B6E")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t3_data = [
        ("Gate Redundancy", "25.0%", "100.0%", "100.0%", "25.0%", "100.0%", "100.0%"),
        ("Physics Violation", "0.0%", "0.0%", "100.0%", "0.0%", "0.0%", "100.0%"),
        ("Semantic Deviation", "0.0%", "27.0%", "100.0%", "25.5%", "100.0%", "100.0%"),
        ("Structural Error", "0.0%", "15.5%", "98.0%", "27.0%", "100.0%", "98.5%"),
        ("Syntax Error", "0.0%", "48.0%", "97.5%", "0.0%", "48.0%", "97.5%")
    ]
    for row_idx, row in enumerate(t3_data):
        for col_idx, text in enumerate(row):
            cell = t3.cell(row_idx + 1, col_idx)
            cell.text = text
            if col_idx in [3, 6]:
                cell.paragraphs[0].runs[0].font.bold = True
            if row_idx % 2 == 1:
                set_cell_background(cell, "F9FAFB")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    fig_heat = os.path.join(FIGURES_DIR, "heatmap_bugtype_vs_system.png")
    if os.path.exists(fig_heat):
        p_fig2 = doc.add_paragraph()
        p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(fig_heat, width=Inches(5.5))
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c2 = p_cap2.add_run("Fig. 2. Defect category repair performance heatmap comparing PyZX, QBugLM, and our Physics-Constrained Agent.")
        r_c2.font.italic = True
        r_c2.font.size = Pt(9.5)

    # Table IV: Statistical Significance
    p_t4 = doc.add_paragraph()
    r_t4 = p_t4.add_run("TABLE IV: PAIRED WILCOXON SIGNED-RANK TESTS AND EFFECT SIZES (COHEN'S d, N = 1,000)")
    r_t4.bold = True
    r_t4.font.size = Pt(10)

    t4 = doc.add_table(rows=7, cols=6)
    t4.alignment = WD_TABLE_ALIGNMENT.CENTER
    t4_headers = ["Comparison", "Metric", "Ours Mean", "Baseline Mean", "p-value", "Cohen's d"]
    for col_idx, h in enumerate(t4_headers):
        cell = t4.cell(0, col_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "084B6E")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    t4_data = [
        ("Ours vs. QBugLM", "Pass@1", "0.775", "0.381", "< 10^-15", "0.799 (Medium-Large)"),
        ("Ours vs. QBugLM", "Pass@3", "0.991", "0.381", "< 10^-15", "1.250 (Very Large)"),
        ("Ours vs. QBugLM", "Physics Validity", "0.992", "0.696", "< 10^-15", "0.639 (Medium)"),
        ("Ours vs. PyZX", "Pass@1", "0.775", "0.050", "< 10^-15", "1.623 (Huge)"),
        ("Ours vs. PyZX", "Pass@3", "0.991", "0.050", "< 10^-15", "3.992 (Extremely Large)"),
        ("Ours vs. PyZX", "Physics Validity", "0.992", "0.155", "< 10^-15", "2.265 (Huge)")
    ]
    for row_idx, row in enumerate(t4_data):
        for col_idx, text in enumerate(row):
            cell = t4.cell(row_idx + 1, col_idx)
            cell.text = text
            if row_idx % 2 == 1:
                set_cell_background(cell, "F9FAFB")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Additional Figures
    fig_iter = os.path.join(FIGURES_DIR, "iterations_used_comparison.png")
    fig_box = os.path.join(FIGURES_DIR, "physics_validity_boxplot.png")
    if os.path.exists(fig_iter):
        p_fig3 = doc.add_paragraph()
        p_fig3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(fig_iter, width=Inches(5.0))
        p_cap3 = doc.add_paragraph()
        p_cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c3 = p_cap3.add_run("Fig. 3. Distribution of repair iterations across systems (Mean = 1.31 iterations for our agent).")
        r_c3.font.italic = True
        r_c3.font.size = Pt(9.5)

    # SECTION VII & VIII
    add_sec("VI. DISCUSSION & THREATS TO VALIDITY")
    doc.add_paragraph(
        "Hardware-Grounded vs. Unconstrained Repair: The empirical data demonstrates that without physical topology and coherence awareness, "
        "even an advanced LLM fails on 100% of physical coupling violations and yields invalid circuits 30.4% of the time. "
        "By enforcing hard hardware constraints at the validator level and feeding diagnostic errors back to the generator, "
        "our multi-agent system recovers 96.0% of initial attempt failures within 3 iterations."
    )
    doc.add_paragraph(
        "Explanation of Structural Overhead: The negative gate count (-9.5%) and circuit depth (-10.0%) reductions reflect necessary physical adaptations. "
        "When repairing non-nearest-neighbor coupling violations on FakeSherbrooke's heavy-hex layout, the agent synthesizes SWAP networks or decomposed bridge gates. "
        "These operations mathematically increase total gate count and depth, but represent the minimal physically executable solution on real hardware."
    )

    add_sec("VII. CONCLUSION")
    doc.add_paragraph(
        "We presented a physics-constrained AI multi-agent framework for automated quantum circuit debugging and repair. "
        "Across 1,000 benchmark OpenQASM 3.0 circuits, our system achieved 99.1% Pass@3 and 99.2% physics validity, substantially outperforming simulation-only and heuristic baselines (p < 10^-15). "
        "The results establish that quantum program repair must integrate physical compilation and execution constraints into the core debugging loop."
    )

    add_sec("REFERENCES")
    refs = [
        "[1] X. Guo, J. Zhao, and P. Zhao, 'On Repairing Quantum Programs Using ChatGPT,' in Proc. IEEE/ACM 5th Int. Workshop on Quantum Software Engineering (Q-SE), 2024, pp. 9-16.",
        "[2] A. B. B. Pham, H. T. Nguyen, and M. Usman, 'QBugLM: An Agentic Benchmarking Framework for LLM-based Quantum Software Debugging,' in Proc. IEEE Int. Conf. on Quantum Software (QSW), 2026.",
        "[3] R. Duncan, A. Kissinger, S. Perdrix, and J. van de Wetering, 'Graph-theoretic simplification of quantum circuits with the ZX-calculus,' Quantum, vol. 4, p. 279, 2020.",
        "[4] G. Li, Y. Ding, and Y. Xie, 'Quarl: A Learning-Based Quantum Circuit Optimizer,' Proc. ACM Program. Lang., vol. 8, no. OOPSLA1, 2024.",
        "[5] S. Tan, L. Lu, D. Xiang, et al., 'HornBro: Homotopy-like Method for Automated Quantum Program Repair,' Proc. ACM Softw. Eng., vol. 2, no. FSE, pp. 734-756, 2025.",
        "[6] J. O. Ernst, A. Chatterjee, T. Franzmeyer, and A. Kuhn, 'Reinforcement Learning for Quantum Control under Physical Constraints,' in Proc. 42nd Int. Conf. Machine Learning (ICML), 2025.",
        "[7] K. Dwivedi, M. Haghparast, and T. Mikkonen, 'Quantum Software Engineering and Quantum Software Development Lifecycle: A Survey,' Cluster Computing, 2024.",
        "[8] Anonymous, 'Quantum Circuit Repair by Gate Prioritisation (QRep),' in Proc. IEEE Int. Conf. on Software Testing, Verification and Validation (ICST), 2026.",
        "[9] Y. Li, H. Pei, L. Huang, B. Yin, and K.-Y. Cai, 'Automatic Repair of Quantum Programs via Unitary Operation,' ACM TOSEM, vol. 33, no. 6, 2024.",
        "[10] C. Yoshida, Y. Ishimoto, O. Nourry, et al., 'Leveraging Mutation Analysis for LLM-based Repair of Quantum Programs,' in Proc. IEEE SANER, 2026.",
        "[11] S. Cao, Z. Zhang, M. Alghadeer, et al., 'Automating Quantum Computing Laboratory Experiments with an Agent-based AI Framework,' Patterns, 2025.",
        "[12] A. C. Hoyt, M. Wang, F. Hua, et al., 'QASMTrans: An End-to-End QASM Compilation Framework with Pulse Generation for Near-Term Quantum Devices,' ACM TOCQ, 2026.",
        "[13] IBM Quantum, 'Build noise models,' IBM Quantum Documentation, accessed Sep. 2026. [Online]. Available: https://quantum.cloud.ibm.com/docs/en/guides/build-noise-models"
    ]
    for r in refs:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_after = Pt(3)
        p_r.paragraph_format.left_indent = Inches(0.25)
        p_r.paragraph_format.first_line_indent = Inches(-0.25)
        run_r = p_r.add_run(r)
        run_r.font.size = Pt(9)

    docx_path = os.path.join(RESEARCH_PAPER_DIR, "research_paper.docx")
    doc.save(docx_path)
    print(f"Generated Word Document: {docx_path}")
    return docx_path

def generate_pdf():
    pdf_path = os.path.join(RESEARCH_PAPER_DIR, "research_paper.pdf")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=1, # Center
        textColor=colors.HexColor('#084B6E'),
        spaceAfter=6
    )
    author_style = ParagraphStyle(
        'DocAuthor',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        alignment=1,
        textColor=colors.HexColor('#222222'),
        spaceAfter=2
    )
    affil_style = ParagraphStyle(
        'DocAffil',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9.5,
        leading=13,
        alignment=1,
        textColor=colors.HexColor('#444444'),
        spaceAfter=14
    )
    abstract_style = ParagraphStyle(
        'DocAbstract',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        alignment=4, # Justify
        textColor=colors.HexColor('#222222'),
        spaceAfter=4
    )
    sec_style = ParagraphStyle(
        'DocSec',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#084B6E'),
        spaceBefore=10,
        spaceAfter=3
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        alignment=4,
        textColor=colors.HexColor('#222222'),
        spaceAfter=4
    )
    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#222222')
    )
    table_header = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white
    )
    caption_style = ParagraphStyle(
        'Caption',
        parent=styles['Italic'],
        fontName='Helvetica-Oblique',
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#444444'),
        spaceAfter=6
    )

    story = []

    # Title & Author
    story.append(Paragraph("Physics-Constrained AI Agents for Automated Quantum Circuit Debugging and Repair", title_style))
    story.append(Paragraph("Alok Pandey", author_style))
    story.append(Paragraph("Department of Computer Science and Engineering, Indian Institute of Technology Patna, India", affil_style))

    # Abstract Box
    abs_html = (
        "<b>Abstract—</b> Quantum-circuit debugging is difficult because a syntactically valid or logically plausible repair may still violate "
        "hardware-native gates, coupling connectivity, or noise-aware execution limits. This paper presents a physics-constrained agentic framework that couples "
        "fault localization and minimal candidate generation with compilation and explicit validation. A structured circuit state records diagnostics and repair history; "
        "failed checks provide feedback for subsequent attempts. The reported benchmark comprises 1,000 buggy OpenQASM 3.0 circuits across five bug classes and three compared systems. "
        "The proposed agent achieves Pass@1 of 77.5%, Pass@3 of 99.1%, and hardware physics validity of 99.2%, with a mean of 1.31 repair iterations. "
        "The evaluation grounds circuits against the IBM FakeSherbrooke 127-qubit superconducting processor model. "
        "Statistical tests and effect size analyses confirm that hardware-grounded validation substantially outperforms simulation-only and heuristic baselines.<br/><br/>"
        "<b><i>Keywords—</i></b> quantum software engineering, circuit debugging, AI agents, automated program repair, hardware-aware compilation, superconducting quantum processors."
    )
    abs_table = Table([[Paragraph(abs_html, abstract_style)]], colWidths=[530])
    abs_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F0F4F8')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(abs_table)
    story.append(Spacer(1, 8))

    # SECTION I
    story.append(Paragraph("I. INTRODUCTION", sec_style))
    story.append(Paragraph(
        "Quantum computing systems require precise agreement between high-level computational specifications and physically executable quantum circuits. "
        "A local textual modification can easily resolve a syntax error while erroneously modifying an entangling target, disrupting basis rotation angles, "
        "or generating multi-qubit interactions prohibited by device topology. Consequently, quantum program repair is intrinsically more demanding "
        "than classical code repair: patches cannot be judged solely by syntactic validity or unconstrained simulation. "
        "Prior language-model-based quantum repair tools (such as QBugLM) rely on unconstrained simulation, ignoring coupling graphs, gate infidelities, and coherence decay limits. "
        "Symbolic optimizers (e.g., PyZX) simplify ZX-calculus diagrams but lack defect-targeted repair mechanisms for algorithmic errors. "
        "To address this gap, we present an end-to-end physics-constrained multi-agent framework that integrates fault localization, minimal candidate synthesis, and rigorous hardware validation into a unified feedback loop, "
        "grounded on the 127-qubit IBM FakeSherbrooke superconducting processor.",
        body_style
    ))

    # SECTION II
    story.append(Paragraph("II. PROBLEM FORMULATION & BENCHMARK TAXONOMY", sec_style))
    story.append(Paragraph(
        "Let C_bug be a buggy circuit, O an independent behavioral oracle, and B a physical hardware backend. "
        "Acceptance requires satisfying syntax, structural, semantic, and hardware physical constraints: Accept(C*) = V_syntax(C*) ∧ V_structural(C*) ∧ V_semantic(C*, O) ∧ V_physics(C*, B). "
        "We evaluate 1,000 circuits across four algorithm families (GHZ, QAOA MaxCut, VQE H2, QFT) and five bug categories (200 instances each).",
        body_style
    ))

    # Table I in PDF
    t1_rows = [
        [Paragraph("<b>Bug Category</b>", table_header), Paragraph("<b>Description & Injection Mechanism</b>", table_header), Paragraph("<b>Validation Criteria</b>", table_header)],
        [Paragraph("Syntax Error", table_cell), Paragraph("Malformed QASM 3.0 syntax, omitted semicolons, corrupt keywords, unclosed brackets.", table_cell), Paragraph("OpenQASM parser diagnostics", table_cell)],
        [Paragraph("Structural Error", table_cell), Paragraph("Missing entangling/superposition gates, inverted control/target qubit operands.", table_cell), Paragraph("AST & operand bounds checking", table_cell)],
        [Paragraph("Gate Redundancy", table_cell), Paragraph("Self-canceling duplicate gates (X·X, H·H, identical CX·CX consecutive pairs).", table_cell), Paragraph("Circuit depth & redundancy reduction", table_cell)],
        [Paragraph("Semantic Deviation", table_cell), Paragraph("Inverted rotation angles (-θ), shifted phase angles, basis mutations (Rx ↔ Rz).", table_cell), Paragraph("Fidelity F = |⟨ψ_gt|ψ_rep⟩|² ≥ 0.999", table_cell)],
        [Paragraph("Physics Violation", table_cell), Paragraph("Non-nearest-neighbor coupling violations on FakeSherbrooke, delays exceeding T2.", table_cell), Paragraph("Heavy-hex graph & coherence limits", table_cell)],
    ]
    t1_pdf = Table(t1_rows, colWidths=[100, 280, 150])
    t1_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#084B6E')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(Paragraph("<b>TABLE I: BENCHMARK DEFECT TAXONOMY (N = 1,000 CIRCUITS)</b>", caption_style))
    story.append(t1_pdf)
    story.append(Spacer(1, 6))

    # SECTION III: ARCHITECTURE
    story.append(Paragraph("III. MULTI-AGENT ARCHITECTURE", sec_style))
    story.append(Paragraph(
        "Our framework coordinates three specialized agents: "
        "(1) <i>Diagnostician Agent</i>: Locates defects using error diagnostics, AST inspection, and statevector mismatch; "
        "(2) <i>Repairer Agent</i>: Synthesizes minimal OpenQASM edits, including syntax fixes, angle inversion, and SWAP-based physical routing; "
        "(3) <i>Validator Agent</i>: Validates candidates against mathematical fidelity oracles and FakeSherbrooke physical coupling/coherence rules. "
        "Failed checks yield structured diagnostic feedback that guides subsequent candidate generation (up to k=3 iterations).",
        body_style
    ))

    # SECTION IV: RESULTS
    story.append(Paragraph("IV. EXPERIMENTAL BENCHMARK RESULTS", sec_style))
    story.append(Paragraph(
        "All three systems were evaluated across all 1,000 circuits in the benchmark. "
        "Our Physics-Constrained Agent achieves 77.5% Pass@1, 99.1% Pass@3, and 99.2% physics validity, with a mean of 1.31 repair iterations. "
        "In contrast, unconstrained LLM repair (QBugLM) achieves only 38.1% Pass@3 and 69.6% physics validity, while the PyZX heuristic achieves 5.0% Pass@3.",
        body_style
    ))

    # Table II in PDF
    t2_rows = [
        [Paragraph("<b>System</b>", table_header), Paragraph("<b>Pass@1 (%)</b>", table_header), Paragraph("<b>Pass@3 (%)</b>", table_header), Paragraph("<b>Physics Valid (%)</b>", table_header), Paragraph("<b>Gate Red.</b>", table_header), Paragraph("<b>Depth Red.</b>", table_header), Paragraph("<b>Iters</b>", table_header)],
        [Paragraph("<b>Ours (Physics-Constrained)</b>", table_cell), Paragraph("<b>77.5 ± 41.8</b>", table_cell), Paragraph("<b>99.1 ± 9.4</b>", table_cell), Paragraph("<b>99.2 ± 8.9</b>", table_cell), Paragraph("-9.5 ± 16.8", table_cell), Paragraph("-10.0 ± 17.2", table_cell), Paragraph("1.31 ± 0.62", table_cell)],
        [Paragraph("QBugLM (Unconstrained)", table_cell), Paragraph("38.1 ± 48.6", table_cell), Paragraph("38.1 ± 48.6", table_cell), Paragraph("69.6 ± 46.0", table_cell), Paragraph("-7.8 ± 15.6", table_cell), Paragraph("-8.1 ± 16.0", table_cell), Paragraph("1.00 ± 0.00", table_cell)],
        [Paragraph("PyZX (Heuristic)", table_cell), Paragraph("5.0 ± 21.8", table_cell), Paragraph("5.0 ± 21.8", table_cell), Paragraph("15.5 ± 36.2", table_cell), Paragraph("-8.4 ± 13.8", table_cell), Paragraph("-6.3 ± 11.0", table_cell), Paragraph("1.00 ± 0.00", table_cell)]
    ]
    t2_pdf = Table(t2_rows, colWidths=[130, 65, 65, 75, 65, 65, 65])
    t2_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#084B6E')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#E8F4F8')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(Paragraph("<b>TABLE II: OVERALL REPAIR PERFORMANCE (N = 1,000 CIRCUITS)</b>", caption_style))
    story.append(t2_pdf)
    story.append(Spacer(1, 6))

    # Pass@k Figure
    fig_pass = os.path.join(FIGURES_DIR, "pass_at_k_comparison.png")
    if os.path.exists(fig_pass):
        story.append(Image(fig_pass, width=400, height=180))
        story.append(Paragraph("Fig. 1. Pass@1 and Pass@3 repair success rates across systems.", caption_style))
        story.append(Spacer(1, 4))

    # Table III in PDF: Defect Category Breakdown
    t3_rows = [
        [Paragraph("<b>Bug Category</b>", table_header), Paragraph("<b>PyZX Pass@3</b>", table_header), Paragraph("<b>QBugLM Pass@3</b>", table_header), Paragraph("<b>Ours Pass@3</b>", table_header), Paragraph("<b>PyZX Phys.</b>", table_header), Paragraph("<b>QBugLM Phys.</b>", table_header), Paragraph("<b>Ours Phys.</b>", table_header)],
        [Paragraph("Gate Redundancy", table_cell), Paragraph("25.0%", table_cell), Paragraph("100.0%", table_cell), Paragraph("<b>100.0%</b>", table_cell), Paragraph("25.0%", table_cell), Paragraph("100.0%", table_cell), Paragraph("<b>100.0%</b>", table_cell)],
        [Paragraph("Physics Violation", table_cell), Paragraph("0.0%", table_cell), Paragraph("0.0%", table_cell), Paragraph("<b>100.0%</b>", table_cell), Paragraph("0.0%", table_cell), Paragraph("0.0%", table_cell), Paragraph("<b>100.0%</b>", table_cell)],
        [Paragraph("Semantic Deviation", table_cell), Paragraph("0.0%", table_cell), Paragraph("27.0%", table_cell), Paragraph("<b>100.0%</b>", table_cell), Paragraph("25.5%", table_cell), Paragraph("100.0%", table_cell), Paragraph("<b>100.0%</b>", table_cell)],
        [Paragraph("Structural Error", table_cell), Paragraph("0.0%", table_cell), Paragraph("15.5%", table_cell), Paragraph("<b>98.0%</b>", table_cell), Paragraph("27.0%", table_cell), Paragraph("100.0%", table_cell), Paragraph("<b>98.5%</b>", table_cell)],
        [Paragraph("Syntax Error", table_cell), Paragraph("0.0%", table_cell), Paragraph("48.0%", table_cell), Paragraph("<b>97.5%</b>", table_cell), Paragraph("0.0%", table_cell), Paragraph("48.0%", table_cell), Paragraph("<b>97.5%</b>", table_cell)],
    ]
    t3_pdf = Table(t3_rows, colWidths=[110, 70, 70, 70, 70, 70, 70])
    t3_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#084B6E')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(Paragraph("<b>TABLE III: REPAIR SUCCESS AND PHYSICS VALIDITY BY DEFECT CATEGORY (N=200 PER CATEGORY)</b>", caption_style))
    story.append(t3_pdf)
    story.append(Spacer(1, 6))

    # Heatmap Figure
    fig_heat = os.path.join(FIGURES_DIR, "heatmap_bugtype_vs_system.png")
    if os.path.exists(fig_heat):
        story.append(Image(fig_heat, width=420, height=220))
        story.append(Paragraph("Fig. 2. Defect category performance heatmap comparing PyZX, QBugLM, and our Physics-Constrained Agent.", caption_style))
        story.append(Spacer(1, 4))

    # Table IV: Wilcoxon
    t4_rows = [
        [Paragraph("<b>Comparison</b>", table_header), Paragraph("<b>Metric</b>", table_header), Paragraph("<b>Ours Mean</b>", table_header), Paragraph("<b>Base Mean</b>", table_header), Paragraph("<b>p-value</b>", table_header), Paragraph("<b>Cohen's d</b>", table_header)],
        [Paragraph("Ours vs. QBugLM", table_cell), Paragraph("Pass@1", table_cell), Paragraph("0.775", table_cell), Paragraph("0.381", table_cell), Paragraph("< 10^-15", table_cell), Paragraph("0.799 (Med-Large)", table_cell)],
        [Paragraph("Ours vs. QBugLM", table_cell), Paragraph("Pass@3", table_cell), Paragraph("0.991", table_cell), Paragraph("0.381", table_cell), Paragraph("< 10^-15", table_cell), Paragraph("1.250 (Very Large)", table_cell)],
        [Paragraph("Ours vs. QBugLM", table_cell), Paragraph("Physics Validity", table_cell), Paragraph("0.992", table_cell), Paragraph("0.696", table_cell), Paragraph("< 10^-15", table_cell), Paragraph("0.639 (Medium)", table_cell)],
        [Paragraph("Ours vs. PyZX", table_cell), Paragraph("Pass@1", table_cell), Paragraph("0.775", table_cell), Paragraph("0.050", table_cell), Paragraph("< 10^-15", table_cell), Paragraph("1.623 (Huge)", table_cell)],
        [Paragraph("Ours vs. PyZX", table_cell), Paragraph("Pass@3", table_cell), Paragraph("0.991", table_cell), Paragraph("0.050", table_cell), Paragraph("< 10^-15", table_cell), Paragraph("3.992 (Extr. Large)", table_cell)],
        [Paragraph("Ours vs. PyZX", table_cell), Paragraph("Physics Validity", table_cell), Paragraph("0.992", table_cell), Paragraph("0.155", table_cell), Paragraph("< 10^-15", table_cell), Paragraph("2.265 (Huge)", table_cell)]
    ]
    t4_pdf = Table(t4_rows, colWidths=[100, 80, 75, 75, 90, 110])
    t4_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#084B6E')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(Paragraph("<b>TABLE IV: STATISTICAL SIGNIFICANCE (PAIRED WILCOXON TESTS & COHEN'S d)</b>", caption_style))
    story.append(t4_pdf)
    story.append(Spacer(1, 6))

    # SECTION V: DISCUSSION & CONCLUSION
    story.append(Paragraph("V. DISCUSSION & CONCLUSION", sec_style))
    story.append(Paragraph(
        "The experimental results establish that physical constraint validation is essential for quantum program repair. "
        "Without hardware coupling and coherence awareness, LLM-based repair yields 30.4% physically unexecutable circuits and completely fails on topology errors. "
        "Negative gate reduction (-9.5%) and depth reduction (-10.0%) values reflect necessary physical routing overhead (e.g. SWAP decompositions) "
        "required to satisfy the 127-qubit heavy-hex coupling lattice. "
        "In conclusion, our physics-constrained multi-agent framework successfully pairs semantic fault localization with hardware validation, "
        "achieving 99.1% Pass@3 across 1,000 benchmark circuits.",
        body_style
    ))

    story.append(Paragraph("REFERENCES", sec_style))
    ref_texts = [
        "[1] X. Guo, J. Zhao, and P. Zhao, 'On Repairing Quantum Programs Using ChatGPT,' in Proc. IEEE/ACM Q-SE, 2024.",
        "[2] A. B. B. Pham, H. T. Nguyen, and M. Usman, 'QBugLM: An Agentic Benchmarking Framework for LLM-based Quantum Software Debugging,' in Proc. IEEE QSW, 2026.",
        "[3] R. Duncan, A. Kissinger, S. Perdrix, and J. van de Wetering, 'Graph-theoretic simplification of quantum circuits with the ZX-calculus,' Quantum, 2020.",
        "[4] G. Li, Y. Ding, and Y. Xie, 'Quarl: A Learning-Based Quantum Circuit Optimizer,' Proc. ACM Program. Lang. (OOPSLA1), 2024.",
        "[5] S. Tan, L. Lu, D. Xiang, et al., 'HornBro: Homotopy-like Method for Automated Quantum Program Repair,' Proc. ACM Softw. Eng. (FSE), 2025.",
        "[6] J. O. Ernst, A. Chatterjee, T. Franzmeyer, and A. Kuhn, 'Reinforcement Learning for Quantum Control under Physical Constraints,' in Proc. ICML, 2025.",
        "[7] K. Dwivedi, M. Haghparast, and T. Mikkonen, 'Quantum Software Engineering and Quantum Software Development Lifecycle: A Survey,' Cluster Computing, 2024.",
        "[8] Anonymous, 'Quantum Circuit Repair by Gate Prioritisation (QRep),' in Proc. IEEE ICST, 2026.",
        "[9] IBM Quantum, 'Build noise models,' IBM Quantum Documentation, 2026."
    ]
    for r in ref_texts:
        story.append(Paragraph(r, ParagraphStyle('Ref', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9.5, spaceAfter=2)))

    doc.build(story)
    print(f"Generated PDF Document: {pdf_path}")
    return pdf_path

if __name__ == "__main__":
    docx_file = generate_docx()
    pdf_file = generate_pdf()
    print("Build complete successfully!")
